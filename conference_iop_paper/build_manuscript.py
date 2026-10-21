"""Build the humanised IOP manuscript using the supplied Word template.

Equations are native editable Office Math, not pictures. Word PDF export is a
separate step so that the PDF is rendered by the same application as the DOCX.
"""
from copy import deepcopy
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
OUTPUT = HERE / 'IOP_Anomaly_and_Health_Monitoring.docx'
TEMPLATE = HERE / 'conference_guidelines/IOP-ConfSer-template.docx'
audit = json.loads((HERE / 'evidence/independent_audit.json').read_text(encoding='utf-8'))
references = json.loads((HERE / 'reference_audit.json').read_text(encoding='utf-8'))['references']


def node(tag, *children):
    element = OxmlElement('m:' + tag)
    for item in children:
        element.append(deepcopy(item))
    return element


def mr(text):
    run = node('r')
    t = OxmlElement('m:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run.append(t)
    return run


def expr(*items):
    children = []
    for item in items:
        if isinstance(item, str):
            children.append(mr(item))
        elif item.tag == qn('m:e'):
            children.extend(list(item))
        else:
            children.append(item)
    return node('e', *children)


def sub(base, index):
    return node('sSub', expr(base), node('sub', mr(index)))


def sup(base, power):
    return node('sSup', expr(base), node('sup', mr(power)))


def frac(numerator, denominator):
    return node('f', node('num', *list(expr(*numerator))), node('den', *list(expr(*denominator))))


def summation(items, lower, upper):
    props=node('naryPr')
    character=OxmlElement('m:chr');character.set(qn('m:val'),'∑');props.append(character)
    location=OxmlElement('m:limLoc');location.set(qn('m:val'),'subSup');props.append(location)
    return node('nary',props,node('sub',mr(lower)),node('sup',mr(upper)),expr(*items))


def equations():
    residual = sub('r', 'i,t,j')
    distance = sup(sub('D', 'i,t'), '2')
    raw = sup(sub('h', 'i,t'), 'raw')
    return {
        'scaling': [sub('x̃', 'i,t,j'), ' = ', frac([sub('x', 'i,t,j'), ' − ', sub('μ', 'k,j')], [sub('σ', 'k,j')])],
        'residual': [residual, ' = ', sub('x̃', 'i,t,j'), ' − ', summation([sub('w', 'm'), '[', sub('f', 'm,j'), '(', sub('z', 'i,t'), ') + ', sup(sub('b', 'i,t,j'), '(m)'), ']'],'m=1','4')],
        'mahalanobis': [distance, ' = ', sup(expr('(r̄ − ', sub('μ', 'r'), ')'), 'T'), sup('Σ̂', '−1'), '(r̄ − ', sub('μ', 'r'), '),   ', sub('s', 'M'), ' = ', frac([distance], [distance, ' + ', sub('q', '0.99')])],
        'rule': [sub('u', 'i,t'), ' = ', frac(['1'], ['p']),summation(['|',sub('r','i,t,j'),'|'],'j=1','p'), ' + 0.5 ',summation(['I(|',sub('r','i,t,j'),'| > 3)'],'j=1','p')],
        'fusion': [sub('S', 'i,t'), ' = ', frac([sub('s', 'IF'), ' + ', sub('s', 'M'), ' + ', sub('s', 'AE'), ' + ', sub('s', 'R')], ['4'])],
        'health_raw': [raw, ' = 100 exp[−λ max(', frac([sub('d', 'i,t')], [sub('d', '0')]), ' − 1, 0)]'],
        'health_smooth': [sub('H', 'i,t'), ' = α', raw, ' + (1 − α)', sub('H', 'i,t−1')],
    }


equation_text = {
    'scaling': 'x̃ᵢ,ₜ,ⱼ = (xᵢ,ₜ,ⱼ − μₖ,ⱼ) / σₖ,ⱼ',
    'residual': 'rᵢ,ₜ,ⱼ = x̃ᵢ,ₜ,ⱼ − ∑ₘ wₘ [fₘ,ⱼ(zᵢ,ₜ) + bᵢ,ₜ,ⱼ⁽ᵐ⁾]',
    'mahalanobis': 'D² = (r̄ − μᵣ)ᵀ Σ̂⁻¹ (r̄ − μᵣ); sM = D² / (D² + q₀.₉₉)',
    'rule': 'uᵢ,ₜ = (1/p) ∑ⱼ |rᵢ,ₜ,ⱼ| + 0.5 ∑ⱼ I(|rᵢ,ₜ,ⱼ| > 3)',
    'fusion': 'Sᵢ,ₜ = (sIF + sM + sAE + sR) / 4',
    'health_raw': 'hrawᵢ,ₜ = 100 exp[−λ max(dᵢ,ₜ/d₀ − 1, 0)]',
    'health_smooth': 'Hᵢ,ₜ = α hrawᵢ,ₜ + (1 − α) Hᵢ,ₜ₋₁',
}


def tables():
    data = [[v['subset'], str(v['train_engines']), f"{v['train_rows']:,}", str(v['test_engines']), f"{v['test_rows']:,}"] for v in audit['dataset']]
    data.append(['Total', '709', '160,359', '707', '104,897'])
    compare = [
        ['Accuracy (%)', '98.99', '97.82'], ['Precision', '0.9045', '0.8999'],
        ['Critical recall', '0.8908', '0.6321'], ['F1', '0.8976', '0.7426'],
        ['Healthy false alerts', '242', '181'], ['Missed critical observations', '281', '947'],
        ['Healthy-cycle false-alert rate (%)', '0.4916', '0.3676'],
    ]
    subset = [[s, f"{v['rows']:,}", f"{v['precision']:.4f}", f"{v['recall']:.4f}", f"{v['f1']:.4f}", str(v['false_positive']), str(v['false_negative'])] for s, v in audit['per_subset'].items()]
    health = [[s, f"{v['rul_correlation']:.4f}", f"{v['monotonicity_score']:.4f}", f"{v['median_health_drop']:.2f}"] for s, v in [*audit['health_per_subset'].items(), ('All', audit['health'])]]
    return {
        'dataset': ('Table 1. C-MAPSS observations reconstructed from the local source files.', ['Subset', 'Train engines', 'Train rows', 'Test engines', 'Test rows'], data, ''),
        'comparison': ('Table 2. Retrospective comparison on the same 51,806 binary condition observations.', ['Measure', 'Linear baseline', 'Four detectors'], compare, 'The baseline uses supervised proxy labels; anomaly references are learned from healthy regions.'),
        'subsets': ('Table 3. Anomaly results by official test subset at the frozen threshold and gate.', ['Subset', 'Binary rows', 'Precision', 'Recall', 'F1', 'FP', 'FN'], subset, 'FP: healthy false-alert observations; FN: missed critical observations.'),
        'health': ('Table 4. Health-index evaluation on all official test trajectories.', ['Subset', 'Pooled Spearman ρ', 'Monotonicity pass fraction', 'Median decrease'], health, 'Decrease is in health points, restricted to engines observed in the critical band. Monotonicity uses the tolerance defined in section 4.3.'),
    }


figure_captions = {
    'architecture': 'Figure 1. Monitoring architecture. Both branches share commissioned residuals, while anomaly scoring and health assessment use separate transformations. IF denotes Isolation Forest and LSTM denotes long short-term memory.',
    'detection': 'Figure 2. Retrospective anomaly evidence. (a) Precision–recall curves for the fused and individual saved scores, with AP in parentheses. The cross marks the context-gated deployed decision. Individual curves are diagnostic rankings, not calibrated standalone ablations. (b) Critical-cycle recall and detection of critical-reaching engines.',
    'health': 'Figure 3. All 707 official test health trajectories, grouped by subset. Each line is one observed engine history. The figure uses the frozen health transformation, with no selection of example engines or additional evaluation smoothing.',
}


def add_hyperlink(paragraph, label, target):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); props.append(color)
    run.append(props)
    t = OxmlElement('w:t'); t.text = label; run.append(t)
    hyperlink.append(run); paragraph._p.append(hyperlink)


def main():
    document = Document(TEMPLATE)
    for child in list(document._element.body):
        if child.tag != qn('w:sectPr'):
            document._element.body.remove(child)
    section = document.sections[0]
    # Preserve the supplied template's exact A4 size and margins.
    width = section.page_width - section.left_margin - section.right_margin
    document.core_properties.title = 'Residual-based turbofan condition monitoring through anomaly fusion and continuous health assessment'
    document.core_properties.author = 'Selvakumar Sadurshan'
    document.core_properties.subject = 'Anomaly detection and health monitoring only'
    document.core_properties.keywords = 'turbofan; anomaly detection; health monitoring; C-MAPSS'
    document.core_properties.comments = 'Evidence-based IOP draft. Affiliation and conference-specific page limit remain to be confirmed.'
    for name in ['IOP-CS-BodyText','IOP-CS-BodyNoIndent','IOP-CS-ReferenceText','IOP-CS-TableContent','IOP-CS-TableLeftColumn','IOP-CS-CaptionText']:
        document.styles[name].paragraph_format.line_spacing = 1.0
        document.styles[name].paragraph_format.space_after = Pt(0)
    document.styles['IOP-CS-BodyNoIndent'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    document.styles['IOP-CS-Abstract'].paragraph_format.line_spacing = 1.0
    document.styles['IOP-CS-Abstract'].paragraph_format.space_after = Cm(1.0)
    for name in ['IOP-CS-SectionHead','IOP-CS-SubsectionHeading']:
        document.styles[name].paragraph_format.keep_with_next = True
        document.styles[name].paragraph_format.line_spacing = 1.0
    for header in [section.header, section.footer, section.first_page_header, section.first_page_footer, section.even_page_header, section.even_page_footer]:
        for paragraph in header.paragraphs:
            paragraph.clear()
    settings = document.settings.element
    if settings.find(qn('w:embedTrueTypeFonts')) is None:
        settings.append(OxmlElement('w:embedTrueTypeFonts'))
    equation_nodes = equations()
    table_definitions = tables()
    expanded = []
    source = (HERE / 'manuscript.md').read_text(encoding='utf-8')
    first_body = True
    frontmatter = True
    equation_number = 0
    for block in re.split(r'\n\s*\n', source.strip()):
        block = block.strip()
        if block.startswith('# '):
            document.add_paragraph(block[2:], style='IOP-CS-Title')
            expanded.append(block)
        elif block.startswith('## '):
            document.add_paragraph(block[3:], style='IOP-CS-SectionHead')
            first_body = True; frontmatter = False; expanded.append(block)
        elif block.startswith('### '):
            document.add_paragraph(block[4:], style='IOP-CS-SubsectionHeading')
            first_body = True; expanded.append(block)
        elif block.startswith('{{equation:'):
            key = block[len('{{equation:'):-2]
            equation_number += 1
            p = document.add_paragraph(style='IOP-CS-BodyNoIndent')
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.tab_stops.add_tab_stop(int(width/2), WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            p.add_run('\t')
            math = node('oMath', *(mr(v) if isinstance(v,str) else v for v in equation_nodes[key]))
            p._p.append(math)
            p.add_run(f'\t({equation_number})')
            expanded.append(f'{equation_text[key]}    ({equation_number})')
        elif block.startswith('{{figure:'):
            key = block[len('{{figure:'):-2]
            p = document.add_paragraph(style='IOP-CS-BodyNoIndent')
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(7)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(HERE / f'figures/{key}.png'), width=width)
            p = document.add_paragraph(figure_captions[key], style='IOP-CS-CaptionText')
            p.paragraph_format.space_after = Pt(8)
            expanded.append(f'![{figure_captions[key]}](figures/{key}.png)')
        elif block.startswith('{{table:'):
            key = block[len('{{table:'):-2]
            caption, headings, rows, note = table_definitions[key]
            p = document.add_paragraph(caption, style='IOP-CS-CaptionText')
            p.paragraph_format.space_before = Pt(8);p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            table = document.add_table(rows=1, cols=len(headings))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            if key == 'comparison':
                fractions = [.56,.22,.22]
            elif key == 'health':
                fractions = [.16,.27,.34,.23]
            elif key == 'subsets':
                fractions = [.14,.18,.16,.15,.15,.11,.11]
            else:
                fractions = [.16,.21,.21,.21,.21]
            for col, fraction in zip(table.columns, fractions):
                col.width = int(width * fraction)
            for i, values in enumerate([headings] + rows):
                cells = table.rows[0].cells if i == 0 else table.add_row().cells
                for j,(cell,value) in enumerate(zip(cells,values)):
                    cell.width = int(width * fractions[j])
                    p = cell.paragraphs[0]
                    p.style = 'IOP-CS-TableLeftColumn' if j==0 else 'IOP-CS-TableContent'
                    p.paragraph_format.space_before = Pt(2);p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.keep_with_next = i < len(rows)
                    run=p.add_run(value)
                    if key=='subsets': run.font.size=Pt(9)
                    if i==0: run.bold=True
                    if i==0 or i==len(rows):
                        props = cell._tc.get_or_add_tcPr()
                        borders=OxmlElement('w:tcBorders')
                        border=OxmlElement('w:bottom');border.set(qn('w:val'),'single');border.set(qn('w:sz'),'6')
                        borders.append(border)
                        if i==0:
                            border=OxmlElement('w:top');border.set(qn('w:val'),'single');border.set(qn('w:sz'),'6');borders.append(border)
                        props.append(borders)
                no_split=OxmlElement('w:cantSplit');table.rows[i]._tr.get_or_add_trPr().append(no_split)
            if note:
                p=document.add_paragraph(note,style='IOP-CS-TableFootnote')
                p.paragraph_format.space_after=Pt(6)
            expanded.append(caption+'\n\n'+'\n'.join(['| '+' | '.join(headings)+' |','| '+' | '.join(['---']*len(headings))+' |']+['| '+' | '.join(row)+' |' for row in rows])+ ('\n\n'+note if note else ''))
        elif block=='{{references}}':
            ref_lines=[]
            for ref in references:
                p=document.add_paragraph(style='IOP-CS-ReferenceText')
                p.paragraph_format.left_indent=Cm(.65);p.paragraph_format.first_line_indent=Cm(-.65)
                p.paragraph_format.space_after=Pt(4)
                text=f"[{ref['id']}] {ref['authors']} {ref['year']} {ref['title']} {ref['publication']}. "
                p.add_run(text)
                add_hyperlink(p,ref['url'],ref['url'])
                ref_lines.append(text+ref['url'])
            expanded.append('\n\n'.join(ref_lines))
        elif frontmatter:
            style='IOP-CS-Author' if block=='Selvakumar Sadurshan' else 'IOP-CS-Affiliation'
            if block.startswith('Abstract.'):
                style='IOP-CS-Abstract'
            p=document.add_paragraph(block,style=style)
            expanded.append(block)
        else:
            document.add_paragraph(block,style='IOP-CS-BodyNoIndent' if first_body else 'IOP-CS-BodyText')
            first_body=False;expanded.append(block)
    document.save(OUTPUT)
    (HERE/'Humanized_Manuscript.md').write_text('\n\n'.join(expanded)+'\n',encoding='utf-8')
    # Keep a machine-readable reference check separate from the submission manuscript.
    audit_lines=['# Reference verification record','Checked on 21 October 2026. Every entry has a real publisher, repository or NASA record.']
    for ref in references:
        audit_lines.append(f"## [{ref['id']}] {ref['title']}\n\n{ref['authors']} ({ref['year']}). {ref['publication']}.\n\nReference: {ref['url']}\n\nChecked at: {ref['verified_at']}\n\nSupports: {ref['supports']}\n\nStatus: {ref['status']}"+ ('\n\nNote: '+ref['note'] if ref.get('note') else ''))
    (HERE/'Verified_References.md').write_text('\n\n'.join(audit_lines)+'\n',encoding='utf-8')
    print(json.dumps({'docx':str(OUTPUT),'native_equations':equation_number,'figures':len(figure_captions),'tables':len(table_definitions),'references':len(references),'body_words':len(source.split())},indent=2))


if __name__=='__main__':
    main()
